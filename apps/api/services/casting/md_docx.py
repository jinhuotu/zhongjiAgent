"""把良率分析 Markdown 转成 Word (.docx) 字节。

覆盖本页导出常见语法：标题、段落、引用、无序/有序列表、管道表格、粗体/斜体。
"""

from __future__ import annotations

import io
import re
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_UL_RE = re.compile(r"^[-*+]\s+(.*)$")
_OL_RE = re.compile(r"^(\d+)[.)]\s+(.*)$")
_BQ_RE = re.compile(r"^>\s?(.*)$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
_INLINE_RE = re.compile(
    r"(\*\*\*[^*]+?\*\*\*|\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)"
)


def _set_run_font(run, *, bold: bool = False, italic: bool = False, code: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(10 if code else 11)
    run.font.name = "Consolas" if code else "微软雅黑"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    east = "Consolas" if code else "微软雅黑"
    rFonts.set(qn("w:eastAsia"), east)
    if code:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def _add_inline(paragraph, text: str) -> None:
    raw = text or ""
    if not raw:
        paragraph.add_run("")
        return
    pos = 0
    for m in _INLINE_RE.finditer(raw):
        if m.start() > pos:
            run = paragraph.add_run(raw[pos : m.start()])
            _set_run_font(run)
        token = m.group(1)
        if token.startswith("***") and token.endswith("***"):
            run = paragraph.add_run(token[3:-3])
            _set_run_font(run, bold=True, italic=True)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, italic=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, code=True)
        else:
            run = paragraph.add_run(token)
            _set_run_font(run)
        pos = m.end()
    if pos < len(raw):
        run = paragraph.add_run(raw[pos:])
        _set_run_font(run)


def _split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [c.strip() for c in text.split("|")]


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    if cols <= 0:
        return
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline(p, row[j] if j < len(row) else "")
            for run in p.runs:
                if i == 0:
                    run.bold = True


def _iter_blocks(lines: Iterable[str]):
    buf: list[str] = []
    table_rows: list[list[str]] = []
    mode: str | None = None

    def flush_para():
        nonlocal buf
        if buf:
            yield ("p", " ".join(buf).strip())
            buf = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            yield ("table", table_rows)
            table_rows = []

    for raw in lines:
        line = raw.rstrip("\n")
        stripped = line.strip()

        if mode == "table":
            if _is_table_row(stripped) and not _TABLE_SEP_RE.match(stripped):
                table_rows.append(_split_table_row(stripped))
                continue
            yield from flush_table()
            mode = None
            # fall through to re-handle this line

        if not stripped:
            yield from flush_para()
            continue

        if _is_table_row(stripped):
            yield from flush_para()
            if _TABLE_SEP_RE.match(stripped):
                mode = "table"
                continue
            table_rows = [_split_table_row(stripped)]
            mode = "table"
            continue

        hm = _HEADING_RE.match(stripped)
        if hm:
            yield from flush_para()
            yield ("h", len(hm.group(1)), hm.group(2).strip())
            continue

        bm = _BQ_RE.match(stripped)
        if bm:
            yield from flush_para()
            yield ("quote", bm.group(1).strip())
            continue

        um = _UL_RE.match(stripped)
        if um:
            yield from flush_para()
            yield ("ul", um.group(1).strip())
            continue

        om = _OL_RE.match(stripped)
        if om:
            yield from flush_para()
            yield ("ol", om.group(2).strip())
            continue

        if stripped == "---" or stripped == "***":
            yield from flush_para()
            yield ("hr", None)
            continue

        buf.append(stripped)

    yield from flush_para()
    yield from flush_table()


def markdown_to_docx_bytes(markdown: str) -> bytes:
    """Markdown → .docx 二进制。"""
    text = (markdown or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    if not text:
        doc.add_paragraph("（空文档）")
    else:
        for block in _iter_blocks(text.split("\n")):
            kind = block[0]
            if kind == "h":
                _, level, content = block
                p = doc.add_heading(level=min(max(level, 1), 4))
                p.clear()
                _add_inline(p, content)
            elif kind == "p":
                p = doc.add_paragraph()
                _add_inline(p, block[1])
            elif kind == "quote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(12)
                run = p.add_run(block[1])
                _set_run_font(run, italic=True)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            elif kind == "ul":
                p = doc.add_paragraph(style="List Bullet")
                p.clear()
                _add_inline(p, block[1])
            elif kind == "ol":
                p = doc.add_paragraph(style="List Number")
                p.clear()
                _add_inline(p, block[1])
            elif kind == "table":
                _flush_table(doc, block[1])
            elif kind == "hr":
                p = doc.add_paragraph("─" * 24)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
